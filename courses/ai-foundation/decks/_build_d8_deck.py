#!/usr/bin/env python3
"""
Build the Domain 8 presentation deck (.pptx) — UKAIC AI Foundation.
Navy/white accessible theme, 16:9, speaker notes on every slide.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

NAVY = RGBColor(0x00, 0x2B, 0x5C)
INK = RGBColor(0x0D, 0x1B, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xEE, 0xF3, 0xFA)
OK = RGBColor(0x2E, 0x7D, 0x32)
BAD = RGBColor(0xC6, 0x28, 0x28)
GREY = RGBColor(0x5A, 0x6B, 0x80)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW = prs.slide_width
SH = prs.slide_height


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(sl, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def tb(sl, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
       font="Arial", space_after=6, line_spacing=1.15):
    box = sl.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else text
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = ln
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
    return box


def header(sl, kicker, title):
    rect(sl, 0, 0, SW, Inches(1.25), NAVY)
    tb(sl, Inches(0.55), Inches(0.16), SW - Inches(1.1), Inches(0.32),
       kicker.upper(), size=12, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
    tb(sl, Inches(0.55), Inches(0.46), SW - Inches(1.1), Inches(0.7),
       title, size=26, bold=True, color=WHITE)


def footer(sl, n):
    tb(sl, Inches(0.55), SH - Inches(0.45), Inches(9), Inches(0.3),
       "UKAIC AI Foundation · Domain 8 — Ethics, Bias, Responsible AI & Governance",
       size=10, color=GREY)
    tb(sl, SW - Inches(1.3), SH - Inches(0.45), Inches(0.8), Inches(0.3),
       str(n), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text


def content_slide(kicker, title, bullets, n, note, size=18):
    sl = add_slide()
    header(sl, kicker, title)
    y = Inches(1.7)
    for b in bullets:
        if isinstance(b, tuple):
            txt, lvl = b
        else:
            txt, lvl = b, 0
        box = sl.shapes.add_textbox(Inches(0.7) + Inches(0.4) * lvl, y,
                                    SW - Inches(1.4) - Inches(0.4) * lvl, Inches(0.9))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ("• " if lvl == 0 else "– ") + txt
        p.line_spacing = 1.15
        for r in p.runs:
            r.font.size = Pt(size if lvl == 0 else size - 2)
            r.font.color.rgb = INK if lvl == 0 else RGBColor(0x33, 0x44, 0x59)
            r.font.name = "Arial"
            r.font.bold = False
        y += Inches(0.62) if lvl == 0 else Inches(0.5)
    footer(sl, n)
    notes(sl, note)
    return sl


def table_slide(kicker, title, headers, rows, n, note, widths=None, fsize=14):
    sl = add_slide()
    header(sl, kicker, title)
    nrows = len(rows) + 1
    ncols = len(headers)
    left, top = Inches(0.6), Inches(1.62)
    width = SW - Inches(1.2)
    height = Inches(0.45) * nrows
    shape = sl.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = shape.table
    if widths:
        avail = width
        for i, w in enumerate(widths):
            tbl.columns[i].width = Emu(int(avail * w))
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(fsize); r.font.bold = True
                r.font.color.rgb = WHITE; r.font.name = "Arial"
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else TINT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(fsize); r.font.color.rgb = INK
                    r.font.name = "Arial"
    footer(sl, n)
    notes(sl, note)
    return sl


n = 0

# ---- 1. Title slide ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
rect(sl, 0, Inches(3.05), SW, Inches(0.06), RGBColor(0x4A, 0x90, 0xD9))
tb(sl, Inches(1.0), Inches(1.5), SW - Inches(2.0), Inches(0.5),
   "UK AI COUNCIL · PROFESSIONAL AI CERTIFICATION PATHWAY", size=14, bold=True,
   color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(2.0), SW - Inches(2.0), Inches(1.0),
   "Domain 8", size=44, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(3.35), SW - Inches(2.0), Inches(1.2),
   "Ethics, Bias, Responsible AI & Governance", size=30, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.5),
   "Responsible use, oversight, disclosure, and a professional decision framework",
   size=16, color=RGBColor(0xC9, 0xDD, 0xF2))
tb(sl, Inches(1.0), Inches(5.4), SW - Inches(2.0), Inches(0.4),
   "8% of the examination · 10 of 120 questions · Foundation level · 2 hours 45 minutes",
   size=13, color=RGBColor(0xA9, 0xC4, 0xE8))
notes(sl, "Domain 8 closes the Foundation course. Smallest examination weighting, but it is "
          "the domain that decides whether everything else is used defensibly. Frame the "
          "session with the central principle: AI can assist while humans remain accountable. "
          "Timing: 2h 45m with a break at 1:10.")

# ---- 2. The question this domain answers ----
n += 1
content_slide("Orientation", "The question this domain answers", [
    "Every other domain asked what you can do with AI",
    "This domain asks who answers for it",
    "The central principle: AI can assist while humans remain accountable",
    "Four domains converge here: prompting (4), verification (5), hallucination (6), data (7)",
    "10 questions in the examination — but the domain that makes the other 110 defensible",
], n,
"Open with the framing statement. Emphasise that candidates often find this domain easy to "
"read and easy to get wrong on the examination. Preview the three recurring traps: the "
"attractive-but-wrong answer, over-reaching, and under-reaching.")

# ---- 3. Domain 8 map ----
n += 1
table_slide("Map", "What Domain 8 covers", ["Area", "Focus"], [
    ["8.1 Bias & Fairness", "Sources of bias; duty in people-affecting uses"],
    ["8.2 Transparency, Accountability & Explainability", "Who is accountable; meaningful accounts"],
    ["8.3 Human Oversight", "AI proposes, a person decides; risk increases with stakes"],
    ["8.4 Disclosure", "When AI use must be disclosed"],
    ["8.5 Policy & Governance", "Policy purpose; risk-based regulatory picture"],
    ["8.6 Professional Decision Framework", "Use / verify / escalate / do not use"],
], n,
"Walk the map quickly. Note that 8.2 and 8.3 carry the conceptual weight, 8.4 and 8.5 are "
"the operational face, and 8.6 is where everything is applied. The examination draws across "
"all six areas but the recurring themes are accountability and oversight.",
widths=[0.42, 0.58], fsize=15)

# ---- 4. 8.1 divider ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.7), SW - Inches(2.0), Inches(0.5),
   "AREA 8.1", size=16, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(3.2), SW - Inches(2.0), Inches(1.2),
   "Bias & Fairness", size=40, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.6),
   "Where bias comes from, and why assessing it is a professional duty", size=18,
   color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "Transition to bias. Signal that the examinable content is the three sources of "
          "bias and the duty triggered by people-affecting uses. Everything else in this "
          "block supports those two facts.")

# ---- 5. Three sources of bias ----
n += 1
table_slide("8.1 Bias & Fairness", "Three sources of bias", 
            ["Source", "What it means", "Example"], [
    ["Training data", "The model learns historical patterns, including unfair ones",
     "Screened on ten years of the company's own hires"],
    ["Design choices", "What the system is optimised for, and what that favours",
     "Optimising for 'career continuity' penalises carers"],
    ["Deployment context", "Same system, different stakes or population",
     "Internal triage model reused for public eligibility"],
], n,
"These three are directly examinable. Make candidates name all three. The point about "
"deployment context is the one candidates miss: a system acceptable in one setting can cause "
"harm in another, so context is part of the bias assessment.",
widths=[0.22, 0.40, 0.38], fsize=14)

# ---- 6. The proxy trap ----
n += 1
content_slide("8.1 Bias & Fairness", "The proxy trap: why removing a field is not enough", [
    "The instinctive fix is to remove names and gender markers",
    "Necessary — but not sufficient",
    "Correlated features carry the same signal indirectly: institution, career gaps, phrasing, postcode",
    "These are proxy signals for the characteristic you removed",
    "Fairness must be tested on outcomes, not asserted from the removal of a field",
], n,
"This is the highest-value teaching point in block 8.1. Candidates reach for field removal as "
"the complete answer. Dismantle it: proxies preserve the signal. The examinable consequence "
"is that fairness requires outcome testing.")

# ---- 7. Evidence: resume study ----
n += 1
sl = add_slide()
header(sl, "8.1 Bias & Fairness", "Evidence: the resume study")
rect(sl, Inches(0.6), Inches(1.6), SW - Inches(1.2), Inches(1.15), TINT)
tb(sl, Inches(0.95), Inches(1.78), SW - Inches(1.9), Inches(0.9),
   "34,500+ resumes generated across 54 occupations · older men rated highest even where the "
   "underlying information was identical", size=19, bold=True, color=NAVY)
y = Inches(3.0)
for t in [
    "Guilbeault, Delecourt & Srinivasa Desikan, published in Nature",
    "ChatGPT wrote work histories portraying women as younger and less experienced",
    "Asked to rate those resumes, it gave older men the highest ratings",
    "Same initial information — different outcomes by name and implied age",
    "Researchers: output filters are 'simplistic' and miss nuanced bias such as gendered ageism",
]:
    tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.45), "• " + t, size=16, color=INK)
    y += Inches(0.5)
tb(sl, Inches(0.8), SH - Inches(0.95), SW - Inches(1.6), Inches(0.35),
   "Source: Stanford Report, 'Researchers uncover AI bias against older working women', 17 October 2025",
   size=11, color=GREY)
footer(sl, n)
notes(sl, "Use this as the evidential anchor for the bias block. Stress that bias is not a "
          "hypothetical concern; it is documented in peer-reviewed research at scale. The "
          "researchers' criticism of output filters is directly relevant: filtering is not a "
          "fix, the bias has to be addressed at a fundamental level.")

# ---- 8. Demo 1 ---
n += 1
sl = add_slide()
header(sl, "Demo 1 · 8 minutes", "The name-swap test")
rect(sl, Inches(0.6), Inches(1.6), SW - Inches(1.2), Inches(0.95), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(1.75), SW - Inches(1.9), Inches(0.7),
   "Two identical candidate summaries. One name suggests a younger candidate, one an older "
   "candidate. Rate each for a mid-level management role.", size=17, color=INK)
y = Inches(2.85)
for t in [
    "Run each in a separate session with an approved tool",
    "Compare the language: 'energy', 'trajectory', 'experience', 'over-qualified'",
    "Note the ratings and whether they diverge",
    "Repeat several times — variability means one run proves nothing (syllabus 2.5)",
    "This is a diagnostic demonstration, not a controlled study",
]:
    tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.45), "• " + t, size=16, color=INK)
    y += Inches(0.5)
footer(sl, n)
notes(sl, "Run the demo live if the room allows, otherwise present the prepared results. "
          "State the caveat explicitly: the evidence for bias comes from the published "
          "research, not from this demonstration. The demo builds intuition about how easily "
          "an automated ranking is steered by irrelevant information; the research "
          "establishes the fact. Use an approved tool in line with organisational policy.")

# ---- 9. Discussion 1 ----
n += 1
content_slide("Discussion 1 · 10 minutes", "Where is your organisation exposed?", [
    "Name one real AI or automated use that affects a person's opportunities, access or treatment",
    "Then apply three questions:",
    ("What data did it learn from, and whose past does that encode?", 1),
    ("What was it optimised for, and what does that make more likely?", 1),
    ("Who is affected if it is wrong — and how would anyone find out?", 1),
    "Resist resolving the uncomfortable cases: the skill is the habit of asking",
], n,
"Facilitate in pairs then take two or three examples with the whole room. Most groups surface "
"at least one genuinely uncomfortable case. Do not resolve them into comfortable answers. "
"The examinable skill is the habit of asking the three questions, not arriving at a "
"reassuring verdict.")

# ---- 10. Formative check 8.1 ----
n += 1
sl = add_slide()
header(sl, "Check · 8.1", "Formative question")
rect(sl, Inches(0.6), Inches(1.7), SW - Inches(1.2), Inches(1.5), TINT)
tb(sl, Inches(0.95), Inches(1.95), SW - Inches(1.9), Inches(1.1),
   "A team removes names and gender markers from applications before AI screening.\n"
   "Why is this not proof that the tool is fair?", size=21, bold=True, color=NAVY)
rect(sl, Inches(0.6), Inches(3.6), SW - Inches(1.2), Inches(1.3), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(3.8), SW - Inches(1.9), Inches(1.0),
   "Because correlated features act as proxies, preserving the same signal indirectly.\n"
   "Fairness has to be tested on outcomes.", size=18, color=RGBColor(0x1F, 0x5B, 0x23))
footer(sl, n)
notes(sl, "Take answers before revealing. The expected answer is the proxy point plus the "
          "outcome-testing consequence. If candidates only give the proxy part, prompt for how "
          "fairness would then be established.")

# ---- 11. 8.2 divider ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.7), SW - Inches(2.0), Inches(0.5),
   "AREA 8.2", size=16, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(3.2), SW - Inches(2.0), Inches(1.2),
   "Transparency, Accountability & Explainability", size=36, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.6),
   "The hinge of the whole domain: somebody is accountable", size=18,
   color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "Signal this as the hinge. Accountability is what makes transparency and oversight "
          "requirements rather than virtues. If nobody is accountable there is no reason to "
          "explain or supervise; because somebody is, both become obligations.")

# ---- 12. Accountability principle ----
n += 1
sl = add_slide()
header(sl, "8.2 Accountability", "The principle: 'the AI did it' is not a defence")
rect(sl, Inches(0.6), Inches(1.6), SW - Inches(1.2), Inches(1.3), TINT)
tb(sl, Inches(0.95), Inches(1.85), SW - Inches(1.9), Inches(1.0),
   "An AI system cannot hold legal or professional liability. Accountability rests with the "
   "human or organisation that chose to deploy and use it.", size=20, bold=True, color=NAVY)
y = Inches(3.15)
for t in [
    "It fails because it names a party with no standing to answer",
    "Not because it is impolite — because responsibility cannot attach to a system",
    "Applies to the deploying organisation, not only the tool provider",
    "'The supplier is liable, so we are covered' is also wrong — it does not transfer your duty",
]:
    tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.5), "• " + t, size=17, color=INK)
    y += Inches(0.55)
tb(sl, Inches(0.8), SH - Inches(0.95), SW - Inches(1.6), Inches(0.35),
   "Syllabus 8.2 · reinforced in 8.3 and 8.6", size=12, color=GREY)
footer(sl, n)
notes(sl, "State the principle three times across the session; repetition is what makes it "
          "stick. Handle the supplier-liability misconception here: supplier liability is a "
          "real commercial and regulatory question but does not remove the deploying "
          "organisation's own accountability.")

# ---- 13. Explainability at Foundation level ----
n += 1
sl = add_slide()
header(sl, "8.2 Explainability", "What explainability means at Foundation level")
half = (SW - Inches(1.5)) / 2
rect(sl, Inches(0.6), Inches(1.65), half, Inches(4.4), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.85), Inches(1.85), half - Inches(0.5), Inches(0.5),
   "IT IS", size=15, bold=True, color=OK)
y = Inches(2.45)
for t in ["Being able to give a meaningful account of how a decision was reached",
          "Stating what the decision was based on",
          "Enough for the decision to be justified and reviewed",
          "A professional-literacy standard"]:
    tb(sl, Inches(0.85), y, half - Inches(0.5), Inches(0.7), "• " + t, size=15, color=INK)
    y += Inches(0.72)
rect(sl, Inches(0.6) + half + Inches(0.3), Inches(1.65), half, Inches(4.4),
     RGBColor(0xFB, 0xE3, 0xE3))
tb(sl, Inches(0.85) + half + Inches(0.3), Inches(1.85), half - Inches(0.5), Inches(0.5),
   "IT IS NOT", size=15, bold=True, color=BAD)
y = Inches(2.45)
for t in ["Publishing the full source code",
          "Revealing internal model weights",
          "Guaranteeing perfect reconstruction years later",
          "A technical or legal standard"]:
    tb(sl, Inches(0.85) + half + Inches(0.3), y, half - Inches(0.5), Inches(0.7),
       "• " + t, size=15, color=INK)
    y += Inches(0.72)
footer(sl, n)
notes(sl, "Both over-claiming and under-claiming are examinable errors. Candidates should be "
          "able to say what explainability is not, because the examination distractor set "
          "includes demands for source code, model weights and perfect reconstruction. A "
          "confidence score is not an explanation and a feature weight is not a reason.")

# ---- 14. Demo 2 ----
n += 1
sl = add_slide()
header(sl, "Demo 2 · 8 minutes", "The unexplainable decision")
rect(sl, Inches(0.6), Inches(1.6), SW - Inches(1.2), Inches(0.9), RGBColor(0xEE, 0xF3, 0xFA))
tb(sl, Inches(0.95), Inches(1.75), SW - Inches(1.9), Inches(0.65),
   "Given: 'Your application was declined by our automated system.'", size=19, bold=True,
   color=NAVY)
y = Inches(2.8)
for t in [
    "In pairs, draft the 'meaningful account' you would give the affected person",
    "Then compare against the technical detail a data scientist might produce:",
    ("Model version, feature weights, confidence score", 1),
    "Which account discharges the professional duty — and why?",
    "Notice: a confidence score is not an explanation; a feature weight is not a reason",
]:
    if isinstance(t, tuple):
        tb(sl, Inches(1.2), y, SW - Inches(2.0), Inches(0.45), "– " + t[0], size=15, color=INK)
    else:
        tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.45), "• " + t, size=16, color=INK)
    y += Inches(0.55)
footer(sl, n)
notes(sl, "The two accounts serve different audiences and only one discharges the "
          "professional duty. Draw out that technical detail is often presented as though it "
          "were an explanation. Candidates should leave able to distinguish the two.")

# ---- 15. Discussion 2 ----
n += 1
sl = add_slide()
header(sl, "Discussion 2 · 10 minutes", "The accountability chain")
rect(sl, Inches(0.6), Inches(1.6), SW - Inches(1.2), Inches(1.15), TINT)
tb(sl, Inches(0.95), Inches(1.78), SW - Inches(1.9), Inches(0.9),
   "A bank deploys a third-party credit scoring model. An applicant is declined.\n"
   "The bank blames the model. The vendor says it is 'advisory only' and the bank decides.",
   size=17, color=INK)
y = Inches(3.0)
for t in ["Who is accountable to the applicant, and why?",
          "Does the vendor's 'advisory only' defence change the answer?",
          "What would the bank need to show in order to answer for the decision?"]:
    tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.5), "• " + t, size=18, color=INK)
    y += Inches(0.6)
rect(sl, Inches(0.6), Inches(5.0), SW - Inches(1.2), Inches(1.1), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(5.18), SW - Inches(1.9), Inches(0.8),
   "Landing point: the bank remains accountable. Advisory status does not transfer "
   "responsibility — the bank chose to act on the advice, and cannot delegate accountability "
   "to a system.", size=16, color=RGBColor(0x1F, 0x5B, 0x23))
footer(sl, n)
notes(sl, "The most valuable discussion in the domain. Let candidates feel the discomfort of "
          "the 'we just use the tool' position before dismantling it. The vendor's position is "
          "commercially real but does not answer the applicant's question: who decided?")

# ---- 16. Check 8.2 ----
n += 1
sl = add_slide()
header(sl, "Check · 8.2", "Formative question")
rect(sl, Inches(0.6), Inches(1.7), SW - Inches(1.2), Inches(1.3), TINT)
tb(sl, Inches(0.95), Inches(1.95), SW - Inches(1.9), Inches(0.9),
   "A manager says the AI made the decision, so she cannot answer for it.\nWhat is wrong with this?",
   size=21, bold=True, color=NAVY)
rect(sl, Inches(0.6), Inches(3.4), SW - Inches(1.2), Inches(1.4), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(3.6), SW - Inches(1.9), Inches(1.1),
   "Accountability requires a party that can answer for the outcome. A system cannot, so "
   "responsibility remains with the human and the organisation.", size=18,
   color=RGBColor(0x1F, 0x5B, 0x23))
footer(sl, n)
notes(sl, "Quick check before the break. Take answers, then reveal. Watch for candidates "
          "reaching for supplier liability — redirect to the deploying organisation.")

# ---- 17. Break ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.9), SW - Inches(2.0), Inches(1.0),
   "BREAK — 10 MINUTES", size=44, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.2), SW - Inches(2.0), Inches(0.6),
   "We return to human oversight: AI proposes, a person decides", size=18,
   color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "Break at 1:10. Announce clearly and hold the restart time. On return, do not "
          "re-teach; move straight into 8.3.")

# ---- 18. 8.3 divider ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.7), SW - Inches(2.0), Inches(0.5),
   "AREA 8.3", size=16, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(3.2), SW - Inches(2.0), Inches(1.2),
   "Human Oversight", size=40, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.6),
   "AI proposes while a person decides and owns the outcome", size=18,
   color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "The core principle of this area. Candidates should be able to state it and to "
          "diagnose its inversion in a scenario.")

# ---- 19. Where risk increases ----
n += 1
sl = add_slide()
header(sl, "8.3 Oversight", "Where automated decision-making increases risk")
rect(sl, Inches(0.6), Inches(1.65), SW - Inches(1.2), Inches(1.15), RGBColor(0xFB, 0xE3, 0xE3))
tb(sl, Inches(0.95), Inches(1.85), SW - Inches(1.9), Inches(0.85),
   "Risk increases where the decision materially affects a person's rights, opportunities "
   "or safety.", size=21, bold=True, color=RGBColor(0x8C, 0x1C, 0x1C))
y = Inches(3.1)
for t in ["The criterion is the stakes for the individual — not the technology",
          "Not the cost of the system, and not whether the data was internally collected",
          "Hiring, lending, eligibility, clinical triage, performance decisions all qualify",
          "Higher stakes require oversight proportionate to the risk"]:
    tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.5), "• " + t, size=17, color=INK)
    y += Inches(0.58)
footer(sl, n)
notes(sl, "The testable criterion is stakes for the individual. Drill this, because "
          "distractors offer cost, speed, data origin and internal-versus-external as though "
          "they were the criterion. They are not.")

# ---- 20. Meaningful oversight ----
n += 1
sl = add_slide()
header(sl, "8.3 Oversight", "What meaningful oversight actually requires")
tb(sl, Inches(0.7), Inches(1.6), SW - Inches(1.4), Inches(0.4),
   "The overseer must have all four. Remove any one and oversight is nominal.", size=17,
   color=NAVY, bold=True)
boxes = [("Information", "Sees what the system did and why"),
         ("Authority", "Can disagree and have it stick"),
         ("Time", "Has the space to actually review"),
         ("Competence", "Knows enough to spot the error")]
bw = (SW - Inches(1.7)) / 4
for i, (h, d) in enumerate(boxes):
    x = Inches(0.7) + (bw + Inches(0.1)) * i
    rect(sl, x, Inches(2.2), bw, Inches(2.1), TINT)
    tb(sl, x + Inches(0.2), Inches(2.42), bw - Inches(0.4), Inches(0.5), h, size=18,
       bold=True, color=NAVY)
    tb(sl, x + Inches(0.2), Inches(3.0), bw - Inches(0.4), Inches(1.2), d, size=14, color=INK)
rect(sl, Inches(0.7), Inches(4.6), SW - Inches(1.4), Inches(1.5), RGBColor(0xFB, 0xE3, 0xE3))
tb(sl, Inches(1.0), Inches(4.8), SW - Inches(2.0), Inches(1.2),
   "What it is NOT: a person clicking 'approve' 400 times a day. Post-hoc review that may "
   "never happen. Re-doing every calculation the system performed.", size=16,
   color=RGBColor(0x8C, 0x1C, 0x1C))
footer(sl, n)
notes(sl, "Write the four words on the board. This framing is the most durable thing "
          "candidates can take from 8.3. The 'not' box distinguishes genuine oversight from "
          "accountability theatre, which is the failure mode organisations most often "
          "mistake for a control.")

# ---- 21. Demo 3 automation bias ----
n += 1
sl = add_slide()
header(sl, "Demo 3 · 8 minutes", "Automation bias")
rect(sl, Inches(0.6), Inches(1.6), SW - Inches(1.2), Inches(0.95), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(1.75), SW - Inches(1.9), Inches(0.7),
   "Review a short AI-suggested answer containing one error that requires domain knowledge "
   "to catch. Decide in 3 minutes whether to accept it.", size=17, color=INK)
y = Inches(2.85)
for t in ["Note how many candidates accept the output",
          "Then reveal the error",
          "Automation bias: accepting an automated suggestion because it is fluent and presented as an answer",
          "Fails most easily under time pressure",
          "Links directly to Domain 6: fluency is not accuracy",
          "Oversight fails here not because the human was absent, but because the human was persuaded"]:
    tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.45), "• " + t, size=15, color=INK)
    y += Inches(0.52)
footer(sl, n)
notes(sl, "Time-limit tightly; the pressure is part of the demonstration. Draw out that "
          "critical evaluation is a trained skill, not a disposition. This connects 8.3 back "
          "to Domain 6 and forward to the practical verification habit in Domain 5.")

# ---- 22. Discussion 3 ----
n += 1
table_slide("Discussion 3 · 10 minutes", "Where should the human sit?",
            ["Use", "Oversight reasoning"], [
    ["Automatic approval of expense claims under £500",
     "Low stakes — sampling plus exception handling may suffice. Justify the threshold."],
    ["AI-assisted shortlisting of job applicants",
     "People-affecting — human decides, AI proposes. Bias assessment applies (8.1)."],
    ["AI triaging patients in a clinical setting",
     "Highest stakes — mandatory verification (6.5). Clinician decides every case."],
], n,
"Facilitate as a whole-room discussion. The point is not that oversight must be identical "
"everywhere; it is that oversight should be proportionate to the stakes and candidates must "
"be able to justify the proportionality rather than assume the tool has settled it.",
widths=[0.42, 0.58], fsize=15)

# ---- 23. Check 8.3 ----
n += 1
sl = add_slide()
header(sl, "Check · 8.3", "Formative question")
rect(sl, Inches(0.6), Inches(1.7), SW - Inches(1.2), Inches(1.3), TINT)
tb(sl, Inches(0.95), Inches(1.95), SW - Inches(1.9), Inches(0.9),
   "A system approves expense claims automatically up to £500 and reports no problems.\n"
   "Why is caution warranted?", size=21, bold=True, color=NAVY)
rect(sl, Inches(0.6), Inches(3.4), SW - Inches(1.2), Inches(1.4), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(3.6), SW - Inches(1.9), Inches(1.1),
   "Without oversight there is no mechanism to detect systematic error or unfairness before "
   "it affects many people. 'No reported problems' is not evidence of no problems.",
   size=17, color=RGBColor(0x1F, 0x5B, 0x23))
footer(sl, n)
notes(sl, "The second sentence is the transferable insight. Systems fail quietly and "
          "consistently; oversight is what converts an unseen systematic error into a "
          "visible one.")

# ---- 24. 8.4 divider ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.7), SW - Inches(2.0), Inches(0.5),
   "AREA 8.4", size=16, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(3.2), SW - Inches(2.0), Inches(1.2),
   "Disclosure", size=40, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.6),
   "When must AI use be disclosed?", size=18, color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "Transition to disclosure. This is the operational face of transparency. The "
          "examinable content is the three triggers and the misleading test.")

# ---- 25. Three triggers ----
n += 1
content_slide("8.4 Disclosure", "The three disclosure triggers", [
    "Where required — by law, regulation, contract or organisational policy",
    "Where AI substantially generated the content — the AI produced the substance, not just assisted",
    "Where non-disclosure would mislead — a reasonable person would form a different view",
    "The misleading test handles the cases the first two triggers miss",
    "Not required for assistive uses that do not substantially generate or alter meaning",
], n,
"Present the triggers as a decision list. The misleading test is the one to teach in depth, "
"because it covers a customer believing they speak to a human and a reader assuming human "
"authorship. Give contrast cases: spelling correction, formatting, looking up a source, "
"checking arithmetic are not disclosure triggers.")

# ---- 26. Separate duties ----
n += 1
sl = add_slide()
header(sl, "8.4 Disclosure", "Verification and disclosure are separate duties")
half = (SW - Inches(1.5)) / 2
rect(sl, Inches(0.6), Inches(1.7), half, Inches(2.6), TINT)
tb(sl, Inches(0.9), Inches(1.95), half - Inches(0.6), Inches(0.5), "VERIFICATION",
   size=16, bold=True, color=NAVY)
tb(sl, Inches(0.9), Inches(2.55), half - Inches(0.6), Inches(1.6),
   "Answers: is it true?\n\nFacts, figures, dates, names, quotes and citations checked at "
   "source (5.5, 6.3)", size=16, color=INK)
rect(sl, Inches(0.6) + half + Inches(0.3), Inches(1.7), half, Inches(2.6), TINT)
tb(sl, Inches(0.9) + half + Inches(0.3), Inches(1.95), half - Inches(0.6), Inches(0.5),
   "DISCLOSURE", size=16, bold=True, color=NAVY)
tb(sl, Inches(0.9) + half + Inches(0.3), Inches(2.55), half - Inches(0.6), Inches(1.6),
   "Answers: how was it produced?\n\nWhether AI use must be revealed to the audience (8.4)",
   size=16, color=INK)
rect(sl, Inches(0.6), Inches(4.5), SW - Inches(1.2), Inches(1.6), RGBColor(0xFB, 0xE3, 0xE3))
tb(sl, Inches(0.95), Inches(4.72), SW - Inches(1.9), Inches(1.3),
   "Satisfying one does NOT discharge the other. A fact-checked report can still require "
   "disclosure; a disclosed draft can still be wrong.", size=18,
   color=RGBColor(0x8C, 0x1C, 0x1C), bold=True)
footer(sl, n)
notes(sl, "A very common candidate error: believing verification removes the disclosure "
          "question. The two answer different questions. Use the boxed statement as the "
          "memorable formulation.")

# ---- 27. Disclosure in law ----
n += 1
sl = add_slide()
header(sl, "8.4 Disclosure", "The regulatory picture — and why dates are a trap")
y = Inches(1.7)
for t in [
    "EU AI Act Article 50 transparency obligations apply from 2 August 2026",
    "They are NOT limited to high-risk systems — they cover four situations:",
    ("Direct interaction with people · synthetic content · emotion recognition or biometric "
     "categorisation · deepfakes and public-interest text", 1),
    "Providers must mark synthetic outputs in machine-readable form",
    "High-risk (Annex III) obligations were postponed to 2 December 2027 by the Digital Omnibus",
    "Lesson: know the shape (risk-based, broadening), not a remembered date",
]:
    if isinstance(t, tuple):
        tb(sl, Inches(1.25), y, SW - Inches(2.0), Inches(0.6), "– " + t[0], size=14,
           color=RGBColor(0x33, 0x44, 0x59))
        y += Inches(0.62)
    else:
        tb(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.5), "• " + t, size=16, color=INK)
        y += Inches(0.56)
tb(sl, Inches(0.8), SH - Inches(1.15), SW - Inches(1.6), Inches(0.5),
   "Sources: artificialintelligenceact.eu, Article 50 practical guide, 14 May 2026; "
   "European Commission, 'Safer and more transparent AI', 2 August 2026", size=10, color=GREY)
footer(sl, n)
notes(sl, "Teach the shape, not the date. The postponement is itself the lesson: a candidate "
          "who memorises a compliance date will be wrong within a year. The examinable skill "
          "is recognising that disclosure obligations are risk-based and broadening. Do not "
          "go deeper into legal interpretation — the syllabus excludes it.")

# ---- 28. Disclosure exercise ----
n += 1
table_slide("Exercise · 12 minutes", "Disclosure call: must / consider / not needed",
            ["Case", "Call"], [
    ["Chatbot answers billing queries without identifying itself as AI", "Must disclose"],
    ["Published research summary drafted largely by AI, human-edited", "Consider disclosing"],
    ["Policy document reformatted by AI, meaning unchanged", "Not needed"],
    ["AI-generated image of a real person in a training deck", "Must disclose"],
    ["Email where AI corrected spelling and grammar", "Not needed"],
], n,
"Run as a quick-fire classification with one-sentence justifications. Push on the "
"justification rather than the label. The human-edited summary is deliberately a judgment "
"call: editorial review informs the decision but does not automatically remove the trigger, "
"because the AI produced the substance.",
widths=[0.72, 0.28], fsize=15)

# ---- 29. Check 8.4 ----
n += 1
sl = add_slide()
header(sl, "Check · 8.4", "Formative question")
rect(sl, Inches(0.6), Inches(1.7), SW - Inches(1.2), Inches(1.3), TINT)
tb(sl, Inches(0.95), Inches(1.95), SW - Inches(1.9), Inches(0.9),
   "A draft has been checked for accuracy. Does that remove the need to consider disclosure?",
   size=21, bold=True, color=NAVY)
rect(sl, Inches(0.6), Inches(3.4), SW - Inches(1.2), Inches(1.4), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(3.6), SW - Inches(1.9), Inches(1.1),
   "No. Accuracy and disclosure are separate obligations answering different questions: "
   "is it true, and how was it produced?", size=18, color=RGBColor(0x1F, 0x5B, 0x23))
footer(sl, n)
notes(sl, "Confirm the distinction before moving on. This is one of the two most commonly "
          "missed points in the domain, alongside accountability.")

# ---- 30. 8.5 divider ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.7), SW - Inches(2.0), Inches(0.5),
   "AREA 8.5", size=16, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(3.2), SW - Inches(2.0), Inches(1.2),
   "Policy & Governance", size=40, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.6),
   "Oversight, accountability, review — and a risk-based landscape", size=18,
   color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "Transition to governance. Preview that candidates are not expected to interpret "
          "legislation; they are expected to know the purpose of policy and the shape of the "
          "regulatory landscape.")

# ---- 31. Three elements of governance ----
n += 1
sl = add_slide()
header(sl, "8.5 Governance", "The three elements of governance")
boxes = [("Oversight", "Who watches the use, and how"),
         ("Accountability", "Who answers when it goes wrong, and for what"),
         ("Review", "How policy and systems are revisited as practice changes")]
bw = (SW - Inches(1.6)) / 3
for i, (h, d) in enumerate(boxes):
    x = Inches(0.6) + (bw + Inches(0.2)) * i
    rect(sl, x, Inches(1.7), bw, Inches(2.2), TINT)
    tb(sl, x + Inches(0.25), Inches(1.95), bw - Inches(0.5), Inches(0.6), h, size=20,
       bold=True, color=NAVY)
    tb(sl, x + Inches(0.25), Inches(2.65), bw - Inches(0.5), Inches(1.1), d, size=16,
       color=INK)
rect(sl, Inches(0.6), Inches(4.2), SW - Inches(1.2), Inches(1.8), RGBColor(0xFB, 0xE3, 0xE3))
tb(sl, Inches(0.95), Inches(4.4), SW - Inches(1.9), Inches(1.5),
   "A policy lacking any one element is incomplete.\n"
   "A policy that exists is not necessarily a policy that works — two years predates the "
   "tools staff now use daily.\nSilence in a policy is NOT permission.", size=17,
   color=RGBColor(0x8C, 0x1C, 0x1C))
footer(sl, n)
notes(sl, "The three elements are examinable. The boxed statements are the practical "
          "teaching. The 'silence is not permission' point bridges directly to shadow AI "
          "from Domain 7: a policy gap is a reason to escalate, never a licence to proceed.")

# ---- 32. Regulatory picture ----
n += 1
content_slide("8.5 Governance", "The regulatory picture at awareness level", [
    "Risk-based — obligations scale with the risk of the use",
    "Multi-layered — no single worldwide regime governs all AI use",
    "In motion — requirements are broadening and deadlines have moved",
    "Obligations attach to the USE, and most frameworks bind the DEPLOYER, not only the developer",
    "This is why an organisational policy exists — and why legal interpretation is out of scope",
], n,
"Awareness level only. Candidates are not required to interpret legislation. The three "
"shape statements are the examinable content. The deployer point is the practical insight: an "
"organisation integrating a third-party tool into hiring or clinical workflows becomes the "
"regulated party, so vendor contracts that do not address compliance leave gaps.")

# ---- 33. Discussion 4 ----
n += 1
content_slide("Discussion 4 · 8 minutes", "Audit your own policy", [
    "Does your organisation have an AI policy — and when was it last reviewed?",
    "Does it name who is accountable, or only what is prohibited?",
    "If a colleague wanted to use a tool the policy does not mention, what would they do next?",
    ("And is that the right answer?", 1),
    "Many groups discover their policy is a list of prohibitions with no accountability structure",
], n,
"Facilitate in small groups. That discovery is the learning outcome, so do not soften it. "
"Ask groups to report the gap they found rather than the policy they wish they had.")

# ---- 34. Check 8.5 ----
n += 1
sl = add_slide()
header(sl, "Check · 8.5", "Formative question")
rect(sl, Inches(0.6), Inches(1.7), SW - Inches(1.2), Inches(1.3), TINT)
tb(sl, Inches(0.95), Inches(1.95), SW - Inches(1.9), Inches(0.9),
   "An employee wants to use an unapproved AI tool. The policy does not mention it. What now?",
   size=21, bold=True, color=NAVY)
rect(sl, Inches(0.6), Inches(3.4), SW - Inches(1.2), Inches(1.4), RGBColor(0xE8, 0xF2, 0xE8))
tb(sl, Inches(0.95), Inches(3.6), SW - Inches(1.9), Inches(1.1),
   "Escalate for approval. A policy gap is a reason to escalate, not to proceed on personal "
   "judgement — silence is not permission.", size=18, color=RGBColor(0x1F, 0x5B, 0x23))
footer(sl, n)
notes(sl, "Reinforce the escalation principle. This is the same point tested in Domain 7.3, "
          "so candidates who have read that domain should already have it.")

# ---- 35. 8.6 divider ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(2.7), SW - Inches(2.0), Inches(0.5),
   "AREA 8.6", size=16, bold=True, color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(3.2), SW - Inches(2.0), Inches(1.2),
   "Professional Decision Framework", size=36, bold=True, color=WHITE)
tb(sl, Inches(1.0), Inches(4.5), SW - Inches(2.0), Inches(0.6),
   "Use / use with verification / escalate / do not use", size=18,
   color=RGBColor(0xC9, 0xDD, 0xF2))
footer(sl, n)
notes(sl, "Final area. Everything in the domain converges on this framework. It is the "
          "practical output of the session.")

# ---- 36. The framework ----
n += 1
table_slide("8.6 Framework", "Applying the framework — when unsure, be more cautious",
            ["Outcome", "When it applies"], [
    ["Use", "Routine, low-stakes, approved tool, no personal or confidential data, no people-affecting decision"],
    ["Use with verification", "Output will be acted on, shared or published; verify facts, figures, dates, names, quotes, citations at source"],
    ["Escalate", "High stakes, unclear policy, unfamiliar tool, legal or medical territory, or you are not confident"],
    ["Do not use", "Prohibited use, confidential data with no approved tool, or a decision requiring qualified human expertise"],
], n,
"The four outcomes are directly examinable, along with the 'when unsure, be more cautious' "
"instruction. Emphasise that the framework is not a risk-avoidance device: it is what makes "
"confident, fast use of AI possible, because it identifies the cases requiring verification "
"or escalation and leaves the rest to proceed.",
widths=[0.26, 0.74], fsize=15)

# ---- 37. Consolidation drill ----
n += 1
table_slide("Consolidation · 10 minutes", "The four-outcome sort",
            ["Scenario", "Outcome"], [
    ["Shortening a public-facing document for a general audience", "Use with verification"],
    ["Drafting a first version of a routine internal update", "Use"],
    ["Deciding whether to dismiss an employee for misconduct", "Escalate / do not use"],
    ["Pasting a client contract with personal data into a free public tool", "Do not use"],
    ["Producing a credit score that affects an applicant", "Escalate"],
    ["Asking an approved tool to suggest a synonym", "Use"],
    ["Summarising a board-level contract for a decision", "Use with verification (read the original)"],
    ["Using an unapproved tool because a colleague recommended it", "Escalate"],
], n,
"Run as a fast sort. Push on the justification, not the label. Several scenarios have more "
"than one defensible answer if the reasoning is sound; the examination tests reasoning. Note "
"that the board contract case also engages 5.2 — the original must be read for high-stakes "
"documents.",
widths=[0.66, 0.34], fsize=14)

# ---- 38. Five principles ----
n += 1
sl = add_slide()
header(sl, "Consolidation", "The five principles to take away")
y = Inches(1.75)
for i, t in enumerate([
    "Accountability stays with the human and the organisation",
    "AI proposes; a person decides and owns the outcome",
    "Oversight is proportionate to the stakes",
    "Disclose where non-disclosure would mislead",
    "When unsure, be more cautious — escalate rather than assume",
], 1):
    rect(sl, Inches(0.7), y, Inches(0.55), Inches(0.55), NAVY)
    tb(sl, Inches(0.7), y + Inches(0.09), Inches(0.55), Inches(0.4), str(i), size=20,
       bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(sl, Inches(1.5), y + Inches(0.06), SW - Inches(2.4), Inches(0.5), t, size=19, color=INK)
    y += Inches(0.85)
rect(sl, Inches(0.7), Inches(6.05), SW - Inches(1.4), Inches(0.85), TINT)
tb(sl, Inches(1.0), Inches(6.2), SW - Inches(2.0), Inches(0.6),
   "The central principle, restated: AI can assist while humans remain accountable.",
   size=18, bold=True, color=NAVY)
footer(sl, n)
notes(sl, "Recall the five principles orally and have candidates supply them before "
          "revealing. The final line is the sentence the domain exists to justify; candidates "
          "should be able to explain how each of the four framework outcomes serves it.")

# ---- 39. Exam technique ----
n += 1
table_slide("Exam technique", "Three recurring traps in Domain 8",
            ["Trap", "How it appears", "The correction"], [
    ["The plausible-but-wrong answer",
     "'Cross-check it against another AI' · 'The supplier is liable'",
     "Cross-checking an AI against an AI is not verification (6.4); supplier liability does not remove yours"],
    ["Over-reaching on the requirement",
     "Publish source code · reveal weights · prohibit all automated decisions",
     "The standard is professional literacy, not technical or legal perfection"],
    ["Under-reaching on the duty",
     "Silence in policy means permission · advisory output needs no oversight · no bias assessment",
     "Oversight, escalation and bias assessment are duties, not options"],
], n,
"Close the exam-technique block with the reliable test: ask who is accountable, and could "
"they show their work? Domain 8 is only 10 questions, so each is worth 0.8% of the paper; "
"the traps above account for most errors.",
widths=[0.26, 0.37, 0.37], fsize=13)

# ---- 40. Course close ----
n += 1
sl = add_slide()
rect(sl, 0, 0, SW, SH, NAVY)
tb(sl, Inches(1.0), Inches(1.4), SW - Inches(2.0), Inches(0.6),
   "END OF DOMAIN 8 · END OF THE FOUNDATION SYLLABUS", size=15, bold=True,
   color=RGBColor(0xA9, 0xC4, 0xE8))
tb(sl, Inches(1.0), Inches(2.0), SW - Inches(2.0), Inches(1.0),
   "Next: the examination", size=40, bold=True, color=WHITE)
y = Inches(3.3)
for t in ["Domain 8 quiz — 31 questions, 70% pass standard, complete in your own time",
          "Final examination — 120 questions across all eight domains",
          "120 minutes · closed book · proctored",
          "Pass mark 70% — 84 of 120 questions",
          "No question requires knowledge of a specific commercial product"]:
    tb(sl, Inches(1.0), y, SW - Inches(2.0), Inches(0.5), "• " + t, size=18,
       color=RGBColor(0xE3, 0xEE, 0xFA))
    y += Inches(0.58)
footer(sl, n)
notes(sl, "Close the course. Confirm the examination structure precisely: 120 questions, "
          "120 minutes, closed book, proctored, 70% pass mark (84 of 120). Release the Domain "
          "8 quiz and point candidates to the final examination paper. Remind them the "
          "certification is awarded on the examination alone; training completion does not "
          "confer it.")

out = Path("/home/billy-forrest/Documents/Hermes/UKAIC-AI-Foundation/decks/domain8-ethics-bias-governance-slides.pptx")
prs.save(out)
print(f"WROTE {out}")
print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
